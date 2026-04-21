"""Memory card import endpoint — Excel / Word / Text → CardSet + Cards."""
import io
import json
import re
import uuid

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import select, func, and_
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.core.database import get_db
from app.core.response import ApiResponse
from app.core.security import decode_access_token, oauth2_scheme
from app.models.models import CardSet, MemoryCard

router = APIRouter()
settings = get_settings()


class DeckCreateRequest(BaseModel):
    name: str = Field(..., min_length=1, max_length=200)
    description: str | None = None
    source_type: str = Field(default="manual", max_length=20)


class DeckUpdateRequest(BaseModel):
    name: str | None = Field(None, min_length=1, max_length=200)
    description: str | None = None
    is_pinned: bool | None = None


class ImportConfirmRequest(BaseModel):
    name: str = Field(..., min_length=1, max_length=200)
    source_type: str = Field(default="manual", max_length=20)
    domain: str = Field(default="通用", max_length=50)
    cards: list[dict]


async def get_current_user_id(token: str = Depends(oauth2_scheme)) -> str:
    payload = decode_access_token(token)
    user_id = payload.get("sub")
    if not user_id:
        raise HTTPException(status_code=401, detail="无效的认证")
    return user_id


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _zh_ratio(text: str) -> float:
    compact = text.replace(" ", "")
    if not compact:
        return 0.0
    zh_chars = sum(1 for c in compact if "\u4e00" <= c <= "\u9fff")
    return zh_chars / len(compact)


def _normalize_card(pair: dict, domain: str = "通用", sort_order: int = 0) -> dict:
    return {
        "source_text": _cell_text(pair.get("source_text")),
        "target_text": _cell_text(pair.get("target_text")),
        "source_lang": pair.get("source_lang") or "en",
        "target_lang": pair.get("target_lang") or "zh",
        "domain": pair.get("domain") or domain,
        "difficulty": int(pair.get("difficulty") or 3),
        "card_type": pair.get("card_type") or "bilingual",
        "sort_order": int(pair.get("sort_order") if pair.get("sort_order") is not None else sort_order),
        "extra_data": pair.get("extra_data"),
    }


def _preview_payload(name: str, source_type: str, cards: list[dict]) -> dict:
    return {
        "name": name,
        "source_type": source_type,
        "card_count": len(cards),
        "preview": cards[:5],
        "cards": cards,
    }


def parse_plain_text_cards(text: str, domain: str = "general") -> list[dict]:
    """Deterministic fallback for pasted vocab lists when LLM keys are unavailable."""
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    cards: list[dict] = []
    pending: dict | None = None

    def make_card(left: str, right: str, index: int) -> dict | None:
        left = _cell_text(left)
        right = _cell_text(right)
        if not left or not right or left == right:
            return None
        left_is_zh = _zh_ratio(left) > _zh_ratio(right)
        source_text = right if left_is_zh else left
        target_text = left if left_is_zh else right
        return {
            "source_text": source_text,
            "target_text": target_text,
            "source_lang": "zh" if _zh_ratio(source_text) > 0.3 else "en",
            "target_lang": "zh" if _zh_ratio(target_text) > 0.3 else "en",
            "domain": domain,
            "difficulty": 3,
            "card_type": "term" if max(len(left), len(right)) < 30 else "sentence",
            "sort_order": index,
        }

    for line in lines:
        cleaned = re.sub(r"^\s*(?:[-*]|\d+[.)、])\s*", "", line).strip()
        parts = re.split(r"\s*(?:\t|\||,|，|:|：| - | – | — )\s*", cleaned, maxsplit=1)
        if len(parts) == 2:
            card = make_card(parts[0], parts[1], len(cards))
            if card:
                cards.append(card)
            pending = None
            continue

        lang = "zh" if _zh_ratio(cleaned) > 0.3 else "en"
        if pending and pending["lang"] != lang:
            card = make_card(pending["text"], cleaned, len(cards))
            if card:
                cards.append(card)
            pending = None
        else:
            pending = {"text": cleaned, "lang": lang}

    seen = set()
    unique_cards = []
    for card in cards:
        key = (card["source_text"].casefold(), card["target_text"].casefold())
        if key in seen:
            continue
        seen.add(key)
        card["sort_order"] = len(unique_cards)
        unique_cards.append(card)
    return unique_cards


def _card_set_payload(card_set: CardSet, mastered_count: int = 0) -> dict:
    mastery_rate = round(mastered_count / card_set.card_count * 100, 1) if card_set.card_count else 0
    return {
        "id": card_set.id,
        "name": card_set.name,
        "description": card_set.description,
        "source_type": card_set.source_type,
        "card_count": card_set.card_count,
        "mastered_count": mastered_count,
        "mastery_rate": mastery_rate,
        "is_pinned": card_set.is_pinned,
        "created_at": card_set.created_at.isoformat() if card_set.created_at else None,
    }


# ═══════════════════════════════════════
# CardSet CRUD
# ═══════════════════════════════════════

@router.get("/sets")
async def list_card_sets(
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(CardSet)
        .where(CardSet.user_id == user_id)
        .order_by(CardSet.is_pinned.desc(), CardSet.updated_at.desc())
    )
    sets = result.scalars().all()

    mastered_result = await db.execute(
        select(MemoryCard.card_set_id, func.count(MemoryCard.id))
        .where(and_(MemoryCard.user_id == user_id, MemoryCard.is_mastered == True))
        .group_by(MemoryCard.card_set_id)
    )
    mastered_by_set = {row[0]: row[1] for row in mastered_result.all()}
    return ApiResponse.success(data=[_card_set_payload(s, mastered_by_set.get(s.id, 0)) for s in sets])


@router.get("/decks")
async def list_decks(
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    """Whitepaper v2 deck list endpoint."""
    return await list_card_sets(user_id=user_id, db=db)


@router.post("/decks", status_code=201)
async def create_deck(
    body: DeckCreateRequest,
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    card_set = CardSet(
        id=str(uuid.uuid4()),
        user_id=user_id,
        name=body.name,
        description=body.description,
        source_type=body.source_type,
        card_count=0,
    )
    db.add(card_set)
    await db.flush()
    await db.refresh(card_set)
    return ApiResponse.success(data=_card_set_payload(card_set))


@router.get("/sets/{set_id}")
async def get_card_set(
    set_id: str,
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=500),
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(CardSet).where(CardSet.id == set_id, CardSet.user_id == user_id)
    )
    card_set = result.scalar_one_or_none()
    if not card_set:
        raise HTTPException(status_code=404, detail="卡片集不存在")

    # Get cards
    cards_result = await db.execute(
        select(MemoryCard)
        .where(MemoryCard.card_set_id == set_id)
        .order_by(MemoryCard.sort_order, MemoryCard.created_at)
        .offset((page - 1) * page_size)
        .limit(page_size)
    )
    cards = cards_result.scalars().all()

    total_result = await db.execute(
        select(func.count(MemoryCard.id)).where(MemoryCard.card_set_id == set_id)
    )
    total = total_result.scalar() or 0
    mastered_total = (await db.execute(
        select(func.count(MemoryCard.id)).where(
            and_(MemoryCard.card_set_id == set_id, MemoryCard.is_mastered == True)
        )
    )).scalar() or 0

    return ApiResponse.success(data={
        "id": card_set.id,
        "name": card_set.name,
        "description": card_set.description,
        "source_type": card_set.source_type,
        "card_count": card_set.card_count,
        "mastered_count": mastered_total,
        "mastery_rate": round(mastered_total / card_set.card_count * 100, 1) if card_set.card_count else 0,
        "is_pinned": card_set.is_pinned,
        "created_at": card_set.created_at.isoformat() if card_set.created_at else None,
        "cards": {
            "items": [
                {
                    "id": c.id,
                    "source_text": c.source_text,
                    "target_text": c.target_text,
                    "source_lang": c.source_lang,
                    "target_lang": c.target_lang,
                    "domain": c.domain,
                    "difficulty": c.difficulty,
                    "card_type": c.card_type,
                    "sort_order": c.sort_order,
                    "is_mastered": c.is_mastered,
                    "review_count": c.review_count,
                }
                for c in cards
            ],
            "total": total,
            "page": page,
            "page_size": page_size,
        },
    })


@router.get("/decks/{deck_id}")
async def get_deck(
    deck_id: str,
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=500),
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    """Whitepaper v2 deck detail endpoint."""
    return await get_card_set(deck_id, page=page, page_size=page_size, user_id=user_id, db=db)


@router.put("/decks/{deck_id}")
async def update_deck(
    deck_id: str,
    body: DeckUpdateRequest,
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(CardSet).where(CardSet.id == deck_id, CardSet.user_id == user_id)
    )
    card_set = result.scalar_one_or_none()
    if not card_set:
        raise HTTPException(status_code=404, detail="卡片集不存在")

    for key, value in body.model_dump(exclude_unset=True).items():
        setattr(card_set, key, value)

    await db.flush()
    await db.refresh(card_set)
    return ApiResponse.success(data=_card_set_payload(card_set))


@router.delete("/sets/{set_id}")
async def delete_card_set(
    set_id: str,
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(CardSet).where(CardSet.id == set_id, CardSet.user_id == user_id)
    )
    card_set = result.scalar_one_or_none()
    if not card_set:
        raise HTTPException(status_code=404, detail="卡片集不存在")
    await db.delete(card_set)
    return ApiResponse.success(message="已删除")


@router.delete("/decks/{deck_id}")
async def delete_deck(
    deck_id: str,
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    """Whitepaper v2 deck delete endpoint."""
    return await delete_card_set(deck_id, user_id=user_id, db=db)


# ═══════════════════════════════════════
# File Import — Excel / Word
# ═══════════════════════════════════════

def parse_excel(file_bytes: bytes) -> list[dict]:
    """Parse Excel file: Chinese/English columns, optional category and notes."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    pairs = []
    for row in ws.iter_rows(min_row=1, max_row=2000, values_only=True):
        if not row or len(row) < 2:
            continue
        a, b = _cell_text(row[0]), _cell_text(row[1])
        if not a or not b:
            continue
        header_probe = f"{a}|{b}".lower()
        if any(token in header_probe for token in ["中文", "英文", "source", "target", "原文", "译文"]):
            continue

        # Auto-detect: if A has Chinese chars, A=zh B=en
        if _zh_ratio(a) > _zh_ratio(b):
            source_text, target_text = b, a
            source_lang, target_lang = "en", "zh"
        else:
            source_text, target_text = a, b
            source_lang, target_lang = "zh", "en"

        # Whitepaper Sprint 7: < 30 chars = term, otherwise sentence.
        card_type = "term" if max(len(a), len(b)) < 30 else "sentence"
        domain = _cell_text(row[2]) if len(row) > 2 else ""
        note = _cell_text(row[3]) if len(row) > 3 else ""
        extra_data = {"note": note} if note else None

        pairs.append({
            "source_text": source_text,
            "target_text": target_text,
            "source_lang": source_lang,
            "target_lang": target_lang,
            "card_type": card_type,
            "domain": domain or "通用",
            "extra_data": extra_data,
        })
    wb.close()
    return pairs


def parse_word(file_bytes: bytes) -> list[dict]:
    """Parse Word file: group by sections, then pair continuous ZH/EN paragraphs."""
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))

    paragraphs = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        lang = "zh" if _zh_ratio(text) > 0.3 else "en"
        paragraphs.append({"text": text, "lang": lang})

    # Group into sections by "篇章" headers
    sections = []
    current_section = "默认篇章"
    current_zh = []
    current_en = []

    for p in paragraphs:
        if p["text"].startswith("篇章") or p["text"].startswith("篇"):
            # Save previous section
            if current_section and (current_zh or current_en):
                sections.append({
                    "title": current_section,
                    "zh_paragraphs": current_zh,
                    "en_paragraphs": current_en,
                })
            current_section = p["text"]
            current_zh = []
            current_en = []
        elif p["lang"] == "zh":
            current_zh.append(p["text"])
        else:
            current_en.append(p["text"])

    if current_zh or current_en:
        sections.append({
            "title": current_section,
            "zh_paragraphs": current_zh,
            "en_paragraphs": current_en,
        })

    # Pair paragraphs: match ZH[i] with EN[i]
    pairs = []
    for section in sections:
        title = section["title"]
        zh_list = section["zh_paragraphs"]
        en_list = section["en_paragraphs"]
        pair_count = max(len(zh_list), len(en_list))

        for i in range(pair_count):
            zh_text = zh_list[i] if i < len(zh_list) else ""
            en_text = en_list[i] if i < len(en_list) else ""
            if not zh_text and not en_text:
                continue
            pairs.append({
                "source_text": en_text or zh_text,
                "target_text": zh_text or en_text,
                "source_lang": "en" if en_text else "zh",
                "target_lang": "zh" if zh_text else "en",
                "card_type": "paragraph",
                "domain": title,
                "section": title,
            })

    return pairs


async def _preview_import_file(
    file: UploadFile,
    expected_type: str,
    name: str = "",
) -> dict:
    file_bytes = await file.read()
    filename = file.filename or "import"
    lower_name = filename.lower()

    if expected_type == "excel":
        if not lower_name.endswith((".xlsx", ".xls")):
            raise HTTPException(status_code=400, detail="请上传 Excel 文件（.xlsx 或 .xls）")
        pairs = parse_excel(file_bytes)
        source_type = "excel"
    elif expected_type == "word":
        if not lower_name.endswith((".docx", ".doc")):
            raise HTTPException(status_code=400, detail="请上传 Word 文件（.docx 或 .doc）")
        pairs = parse_word(file_bytes)
        source_type = "word"
    else:
        raise HTTPException(status_code=400, detail="不支持的导入类型")

    if not pairs:
        raise HTTPException(status_code=400, detail="未能从文件中解析出任何卡片")

    set_name = name or filename.rsplit(".", 1)[0]
    cards = [_normalize_card(pair, sort_order=i) for i, pair in enumerate(pairs)]
    return _preview_payload(set_name, source_type, cards)


async def _create_card_set_from_cards(
    db: AsyncSession,
    user_id: str,
    name: str,
    source_type: str,
    cards: list[dict],
    domain: str = "通用",
) -> CardSet:
    if not cards:
        raise HTTPException(status_code=400, detail="没有卡片可导入")

    normalized_cards = [_normalize_card(card, domain=domain, sort_order=i) for i, card in enumerate(cards)]
    valid_cards = [c for c in normalized_cards if c["source_text"] and c["target_text"]]
    if not valid_cards:
        raise HTTPException(status_code=400, detail="没有有效卡片可导入")

    card_set = CardSet(
        id=str(uuid.uuid4()),
        user_id=user_id,
        name=name,
        source_type=source_type,
        card_count=len(valid_cards),
    )
    db.add(card_set)

    for i, card_data in enumerate(valid_cards):
        extra_data = card_data.get("extra_data")
        card = MemoryCard(
            id=str(uuid.uuid4()),
            user_id=user_id,
            card_set_id=card_set.id,
            source_text=card_data["source_text"],
            target_text=card_data["target_text"],
            source_lang=card_data["source_lang"],
            target_lang=card_data["target_lang"],
            domain=card_data["domain"],
            difficulty=card_data["difficulty"],
            card_type=card_data["card_type"],
            sort_order=i,
            extra_data=json.dumps(extra_data, ensure_ascii=False) if extra_data else None,
        )
        db.add(card)

    await db.flush()
    await db.refresh(card_set)
    return card_set


@router.post("/import/excel")
async def preview_excel_import(
    file: UploadFile = File(...),
    name: str = Form(""),
    user_id: str = Depends(get_current_user_id),
):
    """Whitepaper Sprint 7: Excel upload → preview + generated card candidates."""
    return ApiResponse.success(data=await _preview_import_file(file, "excel", name))


@router.post("/import/docx")
async def preview_docx_import(
    file: UploadFile = File(...),
    name: str = Form(""),
    user_id: str = Depends(get_current_user_id),
):
    """Whitepaper Sprint 7: Word upload → preview + generated card candidates."""
    return ApiResponse.success(data=await _preview_import_file(file, "word", name))


@router.post("/import/file")
async def import_from_file(
    file: UploadFile = File(...),
    name: str = Form(""),
    domain: str = Form("通用"),
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    """Import cards from Excel (.xlsx) or Word (.docx) file."""
    filename = file.filename or "import"
    lower_name = filename.lower()

    if lower_name.endswith((".xlsx", ".xls")):
        preview = await _preview_import_file(file, "excel", name)
    elif lower_name.endswith((".docx", ".doc")):
        preview = await _preview_import_file(file, "word", name)
    else:
        raise HTTPException(status_code=400, detail="不支持的文件格式，请上传 .xlsx 或 .docx")

    card_set = await _create_card_set_from_cards(
        db=db,
        user_id=user_id,
        name=preview["name"],
        source_type=preview["source_type"],
        cards=preview["cards"],
        domain=domain,
    )
    return ApiResponse.success(data={
        "set_id": card_set.id,
        "name": card_set.name,
        "source_type": card_set.source_type,
        "card_count": card_set.card_count,
        "preview": preview["preview"],
    })


@router.post("/import/confirm")
async def confirm_import(
    body: ImportConfirmRequest,
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    """Whitepaper Sprint 7: confirm previewed cards and persist them as a deck."""
    card_set = await _create_card_set_from_cards(
        db=db,
        user_id=user_id,
        name=body.name,
        source_type=body.source_type,
        cards=body.cards,
        domain=body.domain,
    )
    return ApiResponse.success(data={
        "set_id": card_set.id,
        "deck_id": card_set.id,
        "name": card_set.name,
        "source_type": card_set.source_type,
        "card_count": card_set.card_count,
    })


# ═══════════════════════════════════════
# Text Import — with streaming LLM
# ═══════════════════════════════════════

class TextImportRequest(BaseModel):
    text: str = Field(..., min_length=1)
    name: str = Field(..., min_length=1)
    domain: str = Field(default="通用")
    source_lang: str = Field(default="auto")
    target_lang: str = Field(default="auto")


@router.post("/import/text")
async def import_from_text(
    body: TextImportRequest,
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    """Import cards from pasted text. Returns preview for confirmation."""
    import httpx

    if not settings.DEEPSEEK_API_KEY:
        cards = parse_plain_text_cards(body.text, domain=body.domain)

        async def fallback_generate():
            yield f"data: {json.dumps({'type': 'done', 'cards': cards, 'count': len(cards), 'fallback': 'local'}, ensure_ascii=False)}\n\n"

        return StreamingResponse(fallback_generate(), media_type="text/event-stream")

    # Auto-detect languages
    zh_chars = sum(1 for c in body.text if '\u4e00' <= c <= '\u9fff')
    total_chars = len(body.text.replace(" ", "").replace("\n", ""))
    has_zh = zh_chars > total_chars * 0.2
    has_en = any(c.isascii() and c.isalpha() for c in body.text)

    src_lang = body.source_lang if body.source_lang != "auto" else ("en" if has_en else "zh")
    tgt_lang = body.target_lang if body.target_lang != "auto" else ("zh" if has_zh else "en")

    prompt = f"""将以下文本拆分为中英对照的记忆卡片对。

要求：
1. 每张卡片包含 source_text（{src_lang}）和 target_text（{tgt_lang}）
2. 按语义单元拆分，不要拆分过碎
3. 翻译要准确自然
4. 以 JSON 数组格式输出，每个元素包含 source_text, target_text, card_type（term/sentence/paragraph）
5. 只输出 JSON，不要其他内容

文本：
{body.text[:8000]}"""

    try:
        async with httpx.AsyncClient(timeout=60) as client:
            resp = await client.post(
                "https://api.deepseek.com/chat/completions",
                headers={
                    "Content-Type": "application/json",
                    "Authorization": f"Bearer {settings.DEEPSEEK_API_KEY}",
                },
                json={
                    "model": "deepseek-chat",
                    "messages": [{"role": "user", "content": prompt}],
                    "temperature": 0.3,
                    "stream": True,
                },
            )
            resp.raise_for_status()

            # Stream response for instant feel
            async def generate():
                full_content = ""
                async for line in resp.aiter_lines():
                    if line.startswith("data: "):
                        data = line[6:]
                        if data == "[DONE]":
                            break
                        try:
                            chunk = json.loads(data)
                            delta = chunk["choices"][0].get("delta", {})
                            content = delta.get("content", "")
                            if content:
                                full_content += content
                                # Send progress events
                                yield f"data: {json.dumps({'type': 'chunk', 'content': content}, ensure_ascii=False)}\n\n"
                        except (json.JSONDecodeError, KeyError, IndexError):
                            continue

                # Parse final result
                try:
                    start = full_content.index("[")
                    end = full_content.rindex("]") + 1
                    cards = json.loads(full_content[start:end])
                except (ValueError, json.JSONDecodeError):
                    cards = parse_plain_text_cards(body.text, domain=body.domain)

                yield f"data: {json.dumps({'type': 'done', 'cards': cards, 'count': len(cards)}, ensure_ascii=False)}\n\n"

            return StreamingResponse(generate(), media_type="text/event-stream")
    except Exception:
        cards = parse_plain_text_cards(body.text, domain=body.domain)

        async def fallback_generate():
            yield f"data: {json.dumps({'type': 'done', 'cards': cards, 'count': len(cards), 'fallback': 'local'}, ensure_ascii=False)}\n\n"

        return StreamingResponse(fallback_generate(), media_type="text/event-stream")


@router.post("/import/text/confirm")
async def confirm_text_import(
    body: dict,
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    """Confirm text import: create card set + cards from preview."""
    name = body.get("name", "导入卡片")
    domain = body.get("domain", "通用")
    cards = body.get("cards", [])
    card_set = await _create_card_set_from_cards(
        db=db,
        user_id=user_id,
        name=name,
        source_type="text",
        cards=cards,
        domain=domain,
    )
    return ApiResponse.success(data={
        "set_id": card_set.id,
        "name": name,
        "card_count": card_set.card_count,
    })


# ═══════════════════════════════════════
# Batch operations
# ═══════════════════════════════════════

class BatchCreateRequest(BaseModel):
    card_set_id: str
    cards: list[dict]


@router.post("/batch")
async def batch_create_cards(
    body: BatchCreateRequest,
    user_id: str = Depends(get_current_user_id),
    db: AsyncSession = Depends(get_db),
):
    """Batch create cards into a set."""
    # Verify set ownership
    result = await db.execute(
        select(CardSet).where(CardSet.id == body.card_set_id, CardSet.user_id == user_id)
    )
    card_set = result.scalar_one_or_none()
    if not card_set:
        raise HTTPException(status_code=404, detail="卡片集不存在")

    for i, card_data in enumerate(body.cards):
        card = MemoryCard(
            id=str(uuid.uuid4()),
            user_id=user_id,
            card_set_id=body.card_set_id,
            source_text=card_data.get("source_text", ""),
            target_text=card_data.get("target_text", ""),
            source_lang=card_data.get("source_lang", "en"),
            target_lang=card_data.get("target_lang", "zh"),
            domain=card_data.get("domain", "通用"),
            card_type=card_data.get("card_type", "bilingual"),
            sort_order=card_set.card_count + i,
        )
        db.add(card)

    card_set.card_count += len(body.cards)
    await db.flush()
    return ApiResponse.success(data={"created": len(body.cards)})
